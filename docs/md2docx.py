#!/usr/bin/env python3
"""Minimal markdown -> docx converter (headings, tables, code, lists, bold, inline code)."""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+?`)")


def add_inline(paragraph, text, size=10.5, base_bold=False):
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, CODE_FONT, size - 0.5, bold=base_bold)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            run = paragraph.add_run(chunk)
            set_run_font(run, BODY_FONT, size, bold=base_bold)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.rows[0].cells[0]
    shade(cell, "F4F5F7")
    cell.paragraphs[0].text = ""
    for index, line in enumerate(lines):
        para = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        run = para.add_run(line if line.strip() else " ")
        set_run_font(run, CODE_FONT, 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for col, text in enumerate(header):
        cell = table.rows[0].cells[col]
        cell.text = ""
        add_inline(cell.paragraphs[0], text, size=9.5, base_bold=True)
        shade(cell, "E8EEF7")
    for row_index, row in enumerate(body, start=1):
        for col in range(len(header)):
            cell = table.rows[row_index].cells[col]
            cell.text = ""
            add_inline(cell.paragraphs[0], row[col] if col < len(row) else "", size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(src, dst, title):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    for section in doc.sections:
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        set_run_font(run, BODY_FONT, 22, bold=True, color=(0x1A, 0x2B, 0x4A))
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    i = 0
    total = len(lines)
    while i < total:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            block = []
            while i < total and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block)
            continue

        if stripped.startswith("|") and i + 1 < total and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            while i < total and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                if not re.match(r"^[\s:|-]+$", raw):
                    rows.append([c.strip() for c in raw.split("|")])
                i += 1
            add_table(doc, rows)
            continue

        if stripped == "---":
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            pPr = para._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "C9D3E3")
            borders.append(bottom)
            pPr.append(borders)
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                para = doc.add_paragraph()
                para.paragraph_format.space_before = Pt(6)
                add_inline(para, text, size=17, base_bold=True)
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4A)
            else:
                h = doc.add_heading(level=min(level - 1, 4))
                h.paragraph_format.space_before = Pt(10)
                h.paragraph_format.space_after = Pt(4)
                h.text = ""
                add_inline(h, text, size={2: 14, 3: 12, 4: 11}.get(level, 11), base_bold=True)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
                    run.font.bold = True
            i += 1
            continue

        if stripped.startswith(">"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.space_after = Pt(4)
            add_inline(para, stripped.lstrip("> ").strip(), size=10)
            for run in para.runs:
                run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            para = doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        if stripped == "":
            i += 1
            continue

        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        add_inline(para, stripped)
        i += 1

    doc.core_properties.title = title
    doc.core_properties.author = "CGMD-Engine / Arena"
    doc.save(dst)
    return dst


if __name__ == "__main__":
    print(convert(sys.argv[1], sys.argv[2], sys.argv[3]))
